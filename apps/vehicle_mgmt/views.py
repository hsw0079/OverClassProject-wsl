from django.contrib.admin.views.decorators import staff_member_required
from django.shortcuts import render
from django.utils.timezone import now

from apps.vehicle_mgmt.models import SchoolVehicle, ExternalVehicle, EntryExitRecord, VisitorAppointment, NonMotorVehicle


@staff_member_required
def stats_dashboard(request):
    """系统首页统计看板"""
    today = now().date()

    today_entry_count = EntryExitRecord.objects.filter(
        record_type='entry', record_time__date=today
    ).count()
    today_exit_count = EntryExitRecord.objects.filter(
        record_type='exit', record_time__date=today
    ).count()

    # 当前在校车辆（按车牌号去重）
    entered_today = set(EntryExitRecord.objects.filter(
        record_type='entry', record_time__date=today
    ).values_list('plate_number', flat=True))
    exited_today = set(EntryExitRecord.objects.filter(
        record_type='exit', record_time__date=today
    ).values_list('plate_number', flat=True))
    on_campus_count = len(entered_today - exited_today)

    pending_visitors = VisitorAppointment.objects.filter(status='pending').count()
    blacklisted_count = SchoolVehicle.objects.filter(is_blacklisted=True).count()
    today_arrived = VisitorAppointment.objects.filter(
        status__in=['arrived', 'left']
    ).filter(expected_time__date=today).count()

    recent_records = EntryExitRecord.objects.order_by('-record_time')[:10]

    context = {
        'today_entry_count': today_entry_count,
        'today_exit_count': today_exit_count,
        'on_campus_count': on_campus_count,
        'pending_visitors': pending_visitors,
        'blacklisted_count': blacklisted_count,
        'today_arrived': today_arrived,
        'recent_records': recent_records,
        'total_vehicles': SchoolVehicle.objects.count(),
        'external_vehicles': ExternalVehicle.objects.count(),
        'non_motor_count': NonMotorVehicle.objects.count(),
        'total_records_today': today_entry_count + today_exit_count,
    }
    return render(request, 'admin/stats_dashboard.html', context)



@staff_member_required
def import_excel(request):
    """批量导入在校车辆 Excel"""
    from django.contrib import messages
    from openpyxl import load_workbook
    from io import BytesIO

    results = {'success': 0, 'skip': 0, 'errors': []}

    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, '请选择一个 Excel 文件')
            return render(request, 'admin/import_excel.html')

        try:
            wb = load_workbook(filename=BytesIO(excel_file.read()))
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头
        except Exception as e:
            messages.error(request, f'文件读取失败：{e}')
            return render(request, 'admin/import_excel.html')

        for row in rows:
            if not row[0]:  # 空行跳过
                continue
            plate = str(row[0]).strip()
            vtype_raw = str(row[1]).strip() if len(row) > 1 and row[1] else ''
            brand = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            color = str(row[3]).strip() if len(row) > 3 and row[3] else ''
            owner = str(row[4]).strip() if len(row) > 4 and row[4] else ''
            phone = str(row[5]).strip() if len(row) > 5 and row[5] else ''

            if not plate or not owner:
                results['errors'].append(f'数据不完整：{row}')
                continue

            # 车辆类型映射
            TYPE_MAP = {
                '校车': 'school_bus', '教职工车辆': 'staff', '教职工': 'staff',
                '公务车辆': 'official', '公务': 'official',
                '访客车辆': 'visitor', '访客': 'visitor',
                '送货车辆': 'delivery', '送货': 'delivery',
            }
            vtype = TYPE_MAP.get(vtype_raw, 'staff')

            if SchoolVehicle.objects.filter(plate_number=plate).exists():
                results['skip'] += 1
                continue

            try:
                SchoolVehicle.objects.create(
                    plate_number=plate, vehicle_type=vtype,
                    brand=brand or None, color=color or None,
                    owner_name=owner, owner_phone=phone or None,
                )
                results['success'] += 1
            except Exception as e:
                results['errors'].append(f'{plate}：{e}')

        messages.success(
            request,
            f'导入完成：成功 {results["success"]} 条，跳过 {results["skip"]} 条（已存在），'
            f'失败 {len(results["errors"])} 条'
        )
        if results['errors']:
            for err in results['errors'][:5]:
                messages.warning(request, err)

    return render(request, 'admin/import_excel.html', {'results': results})


@staff_member_required
def download_template(request):
    """下载 Excel 导入模板"""
    from openpyxl import Workbook
    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = '在校车辆导入模板'
    ws.append(['车牌号（必填）', '车辆类型', '品牌', '颜色', '所属人/部门（必填）', '联系电话'])
    # 示例数据
    ws.append(['京A12345', '教职工车辆', '大众', '白色', '张老师', '13800001111'])
    ws['A2'].comment = None

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="在校车辆导入模板.xlsx"'
    wb.save(response)
    return response